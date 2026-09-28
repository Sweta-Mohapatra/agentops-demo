# Databricks notebook source
# /// script
# [tool.databricks.environment]
# environment_version = "6"
# ///
# DBTITLE 1,Jira Intake Agent — Stage 01 of Agentic Delivery Framework
# MAGIC %md
# MAGIC # 01 — Jira Intake Agent
# MAGIC
# MAGIC **Stage 01: Intake** of the Agentic Delivery Framework.
# MAGIC
# MAGIC This notebook reads a Jira ticket (summary, description, **all comments**, and attachments),
# MAGIC analyses any attached screenshots/files via a Foundation Model, validates bronze sources
# MAGIC against `agentops.bronze`, and produces a structured **requirement spec in YAML**.
# MAGIC
# MAGIC ### Prerequisites
# MAGIC - Databricks secret scope `jira-intake` with keys: `jira_url`, `jira_email`, `jira_api_token`
# MAGIC - Foundation model endpoints: text (`databricks-meta-llama-3-3-70b-instruct`), vision (`databricks-llama-4-maverick`)
# MAGIC - Tables in `agentops.bronze` catalog/schema

# COMMAND ----------

# DBTITLE 1,Install Dependencies
# MAGIC %pip install openai pyyaml pdfplumber python-docx openpyxl --quiet
# MAGIC dbutils.library.restartPython()

# COMMAND ----------

# DBTITLE 1,Configuration & Imports
import requests
import json
import base64
import yaml
import re
from dataclasses import dataclass, field
from typing import Optional

# ---------- Databricks widgets (parameters) ----------
dbutils.widgets.text("ticket_key", "DBXCOE-24", "Jira Ticket Key (e.g. PROJ-123)")
dbutils.widgets.text("secret_scope", "jira-intake", "Databricks Secret Scope")
dbutils.widgets.text("fm_endpoint", "databricks-meta-llama-3-3-70b-instruct", "Foundation Model Endpoint")
dbutils.widgets.text("vision_endpoint", "databricks-llama-4-maverick", "Vision Model Endpoint (for images)")
dbutils.widgets.text("catalog", "agentops", "Target Catalog")
dbutils.widgets.text("bronze_schema", "bronze", "Bronze Schema")
dbutils.widgets.text("gold_schema", "gold", "Gold Schema")

TICKET_KEY    = dbutils.widgets.get("ticket_key").strip()
SECRET_SCOPE  = dbutils.widgets.get("secret_scope")
FM_ENDPOINT   = dbutils.widgets.get("fm_endpoint")
VISION_ENDPOINT = dbutils.widgets.get("vision_endpoint")
CATALOG       = dbutils.widgets.get("catalog")
BRONZE_SCHEMA = dbutils.widgets.get("bronze_schema")
GOLD_SCHEMA   = dbutils.widgets.get("gold_schema")

assert TICKET_KEY, "Please provide a Jira ticket key via the widget."
print(f"Processing ticket: {TICKET_KEY}")

# COMMAND ----------

# DBTITLE 1,Jira API Client
# ---------------------------------------------------------------------------
# Jira REST API helpers  (v3 — Jira Cloud 2025+)
# ---------------------------------------------------------------------------

class JiraClient:
    """Lightweight Jira Cloud REST client using API v3 + token auth."""

    API = "rest/api/3"  # Jira Cloud has retired /rest/api/2/search

    def __init__(self, base_url: str, email: str, api_token: str):
        self.base_url = base_url.rstrip("/")
        self.session = requests.Session()
        self.session.auth = (email, api_token)
        self.session.headers.update({"Accept": "application/json"})

    # -- core helpers --------------------------------------------------------
    def _get(self, path: str, **kwargs) -> dict:
        url = f"{self.base_url}/{self.API}/{path}"
        resp = self.session.get(url, **kwargs)
        resp.raise_for_status()
        return resp.json()

    def _get_binary(self, url: str) -> bytes:
        resp = self.session.get(url)
        resp.raise_for_status()
        return resp.content

    # -- ticket data --------------------------------------------------------
    def get_issue(self, key: str) -> dict:
        """Return the full issue JSON (fields + comments + attachments)."""
        return self._get(f"issue/{key}", params={"expand": "renderedFields"})

    def get_comments(self, key: str) -> list[dict]:
        """Return all comments on a ticket."""
        data = self._get(f"issue/{key}/comment")
        return data.get("comments", [])

    def get_attachments(self, issue: dict) -> list[dict]:
        """Extract attachment metadata from an issue payload."""
        return issue.get("fields", {}).get("attachment", [])

    def download_attachment(self, attachment: dict) -> bytes:
        """Download the raw bytes of an attachment."""
        return self._get_binary(attachment["content"])

    # -- epic / hierarchy ---------------------------------------------------
    def search_jql(self, jql: str, fields: str = "summary,status,issuetype,description,comment,attachment,labels,priority", max_results: int = 50) -> list[dict]:
        """Run a JQL query via the v3 search endpoint and return issues."""
        data = self._get("search/jql", params={
            "jql": jql,
            "maxResults": max_results,
            "fields": fields,
        })
        return data.get("issues", [])

    def get_epic_children(self, epic_key: str) -> list[dict]:
        """Return all child issues (stories, tasks, bugs) under an epic."""
        return self.search_jql(f'"Epic Link" = {epic_key}')

    def is_epic(self, issue: dict) -> bool:
        """Check whether the given issue is an Epic."""
        return (issue.get("fields", {}).get("issuetype") or {}).get("name", "").lower() == "epic"


# Instantiate the client from Databricks secrets
jira = JiraClient(
    base_url  = dbutils.secrets.get(SECRET_SCOPE, "jira_url"),
    email     = dbutils.secrets.get(SECRET_SCOPE, "jira_email"),
    api_token = dbutils.secrets.get(SECRET_SCOPE, "jira_api_token"),
)
print(f"Jira client initialised for {jira.base_url} (API v3)")

# COMMAND ----------

# DBTITLE 1,Read Ticket — Summary, Description, Comments, Attachments (+ Epic Children)
# ---------------------------------------------------------------------------
# Pull everything from the Jira ticket (+ children if it's an Epic)
# ---------------------------------------------------------------------------

def extract_text(body) -> str:
    """Best-effort plain-text extraction from Jira body (ADF or string)."""
    if body is None:
        return ""
    if isinstance(body, str):
        return body
    # Atlassian Document Format (ADF) — walk content nodes recursively
    if isinstance(body, dict) and body.get("type") == "doc":
        parts = []
        for node in body.get("content", []):
            parts.append(_adf_node_text(node))
        return "\n".join(parts)
    return str(body)

def _adf_node_text(node: dict) -> str:
    """Recursively extract text from an ADF node."""
    ntype = node.get("type", "")
    if ntype == "text":
        return node.get("text", "")
    if ntype == "hardBreak":
        return "\n"
    children = node.get("content", [])
    joiner = "\n" if ntype in ("bulletList", "orderedList") else " "
    prefix = "• " if ntype == "listItem" else ""
    return prefix + joiner.join(_adf_node_text(c) for c in children)


def read_issue_full(key: str) -> dict:
    """Read one ticket and return a normalised dict of its content."""
    issue = jira.get_issue(key)
    fields = issue["fields"]
    comments_raw = jira.get_comments(key)
    return {
        "key": key,
        "summary": fields.get("summary", ""),
        "description": extract_text(fields.get("description")),
        "labels": fields.get("labels", []),
        "priority": (fields.get("priority") or {}).get("name", ""),
        "status": (fields.get("status") or {}).get("name", ""),
        "type": (fields.get("issuetype") or {}).get("name", ""),
        "comments": [
            {
                "author": (c.get("author") or {}).get("displayName", "unknown"),
                "created": c.get("created", ""),
                "body": extract_text(c.get("body")),
            }
            for c in comments_raw
        ],
        "attachments": fields.get("attachment", []),
        "_raw_issue": issue,
    }


# ---------- Read the primary ticket ----------
primary = read_issue_full(TICKET_KEY)
print(f"{'='*60}")
print(f"PRIMARY TICKET: {primary['key']}")
print(f"  Summary : {primary['summary']}")
print(f"  Status  : {primary['status']}  |  Priority: {primary['priority']}  |  Type: {primary['type']}")
print(f"  Desc    : {len(primary['description'])} chars")
print(f"  Comments: {len(primary['comments'])}")
print(f"  Attachments: {len(primary['attachments'])}")

# ---------- If it's an Epic, also read all child stories ----------
child_tickets = []
if jira.is_epic(primary["_raw_issue"]):
    children = jira.get_epic_children(TICKET_KEY)
    print(f"\n➡ Epic detected — found {len(children)} child issue(s):")
    for child_issue in children:
        child_key = child_issue["key"]
        child = read_issue_full(child_key)
        child_tickets.append(child)
        print(f"\n  CHILD: {child['key']} [{child['type']}] {child['summary']} — {child['status']}")
        print(f"    Desc: {len(child['description'])} chars  |  Comments: {len(child['comments'])}  |  Attachments: {len(child['attachments'])}")
        if child["description"]:
            print(f"    Description preview: {child['description'][:200]}...")
        for c in child["comments"]:
            print(f"    Comment ({c['author']}): {c['body'][:150]}")

# Consolidate for downstream cells
all_tickets = [primary] + child_tickets
ticket_summary     = primary["summary"]
ticket_description = primary["description"]
ticket_labels      = primary["labels"]
ticket_priority    = primary["priority"]
ticket_status      = primary["status"]
ticket_type        = primary["type"]
comments           = primary["comments"]
attachments        = primary["attachments"]

print(f"\n{'='*60}")
print(f"Total tickets ingested: {len(all_tickets)} (1 primary + {len(child_tickets)} children)")

# COMMAND ----------

# DBTITLE 1,Attachment Understanding (images, PDFs, DOCX, text)
# ---------------------------------------------------------------------------
# Attachment understanding: images, PDFs, DOCX, Excel, text
# ---------------------------------------------------------------------------
import io
from openai import OpenAI

workspace_url = spark.conf.get("spark.databricks.workspaceUrl")

# Token: prefer a service-principal token from secrets (job-safe),
# fall back to interactive notebook context token.
try:
    workspace_token = dbutils.secrets.get(SECRET_SCOPE, "databricks_token")
except Exception:
    workspace_token = (
        dbutils.notebook.entry_point.getDbutils()
        .notebook().getContext().apiToken().get()
    )

client = OpenAI(
    base_url=f"https://{workspace_url}/serving-endpoints",
    api_key=workspace_token,
)

def _summarise_text(filename: str, mime: str, text: str) -> str:
    """Send extracted text to the LLM for a concise summary."""
    resp = client.chat.completions.create(
        model=FM_ENDPOINT,
        messages=[
            {"role": "system", "content": "Summarise what this attached file contains. Mention only what is explicitly present."},
            {"role": "user", "content": f"Filename: {filename}\nMIME: {mime}\n\nContent:\n{text[:12000]}"}
        ],
        max_tokens=400,
    )
    return resp.choices[0].message.content.strip()


def attachment_summary(attachment: dict) -> dict:
    """Describe what an attached image, PDF, DOCX, Excel, or text file shows."""
    filename = attachment["filename"]
    mime = attachment.get("mimeType", "application/octet-stream")
    raw = jira.download_attachment(attachment)
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""

    result = {"filename": filename, "mime_type": mime, "description": "", "text_excerpt": ""}

    # ── Images → vision model ──
    if mime.startswith("image/"):
        b64 = base64.b64encode(raw).decode("utf-8")
        resp = client.chat.completions.create(
            model=VISION_ENDPOINT,
            messages=[{"role": "user", "content": [
                {"type": "text", "text": "Describe this Jira attachment screenshot precisely. Focus only on what is visible: UI elements, metrics, tables, filters, charts, labels, and any business requirement hints."},
                {"type": "image_url", "image_url": {"url": f"data:{mime};base64,{b64}"}}
            ]}],
            max_tokens=500,
        )
        result["description"] = resp.choices[0].message.content.strip()
        return result

    # ── PDF → pdfplumber ──
    if mime == "application/pdf" or ext == "pdf":
        import pdfplumber
        text_parts = []
        with pdfplumber.open(io.BytesIO(raw)) as pdf:
            for page in pdf.pages[:30]:          # cap at 30 pages
                text_parts.append(page.extract_text() or "")
        text = "\n".join(text_parts).strip()
        result["text_excerpt"] = text[:2000]
        if text:
            result["description"] = _summarise_text(filename, mime, text)
        else:
            result["description"] = "PDF contains no extractable text (likely scanned images). Consider OCR."
        return result

    # ── DOCX → python-docx ──
    if ext == "docx" or mime == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        import docx
        doc = docx.Document(io.BytesIO(raw))
        text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        result["text_excerpt"] = text[:2000]
        result["description"] = _summarise_text(filename, mime, text) if text else "Empty DOCX."
        return result

    # ── Excel → openpyxl ──
    if ext in ("xlsx", "xls") or "spreadsheet" in mime:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        lines = []
        for ws in wb.worksheets[:5]:             # cap at 5 sheets
            lines.append(f"Sheet: {ws.title}")
            for row in ws.iter_rows(max_row=50, values_only=True):
                lines.append("\t".join(str(c) if c is not None else "" for c in row))
        text = "\n".join(lines)
        result["text_excerpt"] = text[:2000]
        result["description"] = _summarise_text(filename, mime, text) if text else "Empty spreadsheet."
        return result

    # ── Plain text / code / CSV / JSON / YAML ──
    if mime.startswith("text/") or ext in ("txt", "md", "csv", "json", "yaml", "yml", "sql"):
        text = raw.decode("utf-8", errors="ignore")[:12000]
        result["text_excerpt"] = text[:2000]
        result["description"] = _summarise_text(filename, mime, text)
        return result

    # ── Fallback ──
    result["description"] = f"Binary attachment ({mime}); automatic extraction not supported for this type."
    return result


# Process all attachments (primary + children)
all_attachments = attachments[:]
for ct in child_tickets:
    all_attachments.extend(ct.get("attachments", []))

attachment_descriptions = [attachment_summary(a) for a in all_attachments]
for item in attachment_descriptions:
    print(f"\nAttachment: {item['filename']}")
    print(item['description'][:1000])

print(f"\nTotal attachments processed: {len(attachment_descriptions)}")

# COMMAND ----------

# DBTITLE 1,Validate Existing Bronze Sources
# ---------------------------------------------------------------------------
# Validate which bronze tables actually exist in agentops.bronze
# ---------------------------------------------------------------------------
bronze_df = spark.sql(f"""
SELECT table_name
FROM {CATALOG}.information_schema.tables
WHERE table_schema = '{BRONZE_SCHEMA}'
ORDER BY table_name
""")
existing_bronze_tables = [f"{CATALOG}.{BRONZE_SCHEMA}." + r["table_name"] for r in bronze_df.collect()]
print("Existing bronze tables:")
for t in existing_bronze_tables:
    print(f"  - {t}")

# COMMAND ----------

# DBTITLE 1,Generate YAML Requirement Spec
# ---------------------------------------------------------------------------
# Use an LLM to convert the ticket into the exact YAML schema requested
# ---------------------------------------------------------------------------

def build_ticket_context() -> str:
    """Build full context from the primary ticket + all child tickets."""
    parts = [
        f"PRIMARY TICKET KEY: {TICKET_KEY}",
        f"SUMMARY: {ticket_summary}",
        f"TYPE: {ticket_type}",
        f"DESCRIPTION:\n{ticket_description}",
        f"STATUS: {ticket_status}",
        f"PRIORITY: {ticket_priority}",
        f"LABELS: {', '.join(ticket_labels) if ticket_labels else 'None'}",
        "\nPRIMARY TICKET COMMENTS:",
    ]
    for idx, c in enumerate(comments, start=1):
        parts.append(f"[{idx}] {c['author']} @ {c['created']}:\n{c['body']}")
    if not comments:
        parts.append("- None")

    # Child tickets (from Epic)
    if child_tickets:
        parts.append(f"\n{'='*40}")
        parts.append(f"CHILD STORIES UNDER EPIC ({len(child_tickets)} total):")
        for ct in child_tickets:
            parts.append(f"\n--- {ct['key']} [{ct['type']}]: {ct['summary']} ---")
            parts.append(f"Status: {ct['status']}  |  Priority: {ct['priority']}")
            parts.append(f"Description:\n{ct['description']}")
            if ct["comments"]:
                parts.append("Comments:")
                for idx, c in enumerate(ct["comments"], start=1):
                    parts.append(f"  [{idx}] {c['author']} @ {c['created']}:\n  {c['body']}")

    parts.append("\nATTACHMENT DESCRIPTIONS:")
    if attachment_descriptions:
        for a in attachment_descriptions:
            parts.append(f"- {a['filename']} ({a['mime_type']}): {a['description']}")
    else:
        parts.append("- None")

    parts.append("\nEXISTING BRONZE TABLES:")
    for t in existing_bronze_tables:
        parts.append(f"- {t}")

    return "\n".join(parts)

prompt = f"""
You are converting a Jira ticket into a strict YAML requirement spec for a data platform build workflow.

Rules:
1. Read the ticket summary, description, comments, and attachment descriptions.
2. Use ONLY information explicitly stated in the ticket/comments/attachments.
3. Do NOT invent metric definitions, business rules, filters, dimensions, or grain.
4. If something is unclear, place it in open_questions.
5. bronze_sources must contain ONLY tables from the provided EXISTING BRONZE TABLES list.
6. target_table must be in the form agentops.gold.<name>.
7. Output valid YAML only. No markdown fences. No commentary.
8. Keep the exact shape below.

Required YAML shape:
request_id: {TICKET_KEY}
source_type: jira
report_name: <short name for what's being asked>
grain: <the lowest level of detail one output row represents>
metrics:
  - name: <metric name>
    definition: <plain-English definition, from what's actually stated>
dimensions: [<list>]
filters: [<business rules actually stated, not inferred ones>]
bronze_sources: [<only tables that actually exist in agentops.bronze>]
target_table: agentops.gold.<name>
open_questions:
  - <anything ambiguous or not stated — don't guess at it, list it here instead>

Ticket context:
{build_ticket_context()}
"""

def _clean_yaml(raw: str) -> str:
    """Strip markdown fences and any leading/trailing prose from LLM output."""
    text = raw.strip()
    # Remove markdown fences
    if text.startswith("```"):
        text = re.sub(r"^```(?:ya?ml)?\s*\n", "", text)
        text = re.sub(r"\n```\s*$", "", text)
    # If the LLM prepended prose before the YAML, find first 'request_id:'
    idx = text.find("request_id:")
    if idx > 0:
        text = text[idx:]
    return text.strip()


def _generate_yaml(attempt: int = 1) -> tuple[str, dict]:
    """Call the LLM and parse the result; retry up to 3 times on bad YAML."""
    MAX_RETRIES = 3
    last_error = None
    for i in range(MAX_RETRIES):
        temperature = 0.0 if i > 0 else 0.3   # tighter on retries
        resp = client.chat.completions.create(
            model=FM_ENDPOINT,
            messages=[
                {"role": "system", "content": "You produce only strict YAML. No prose, no commentary, no markdown fences."},
                {"role": "user", "content": prompt}
            ],
            max_tokens=1200,
            temperature=temperature,
        )
        raw = resp.choices[0].message.content.strip()
        cleaned = _clean_yaml(raw)
        try:
            parsed = yaml.safe_load(cleaned)
            if isinstance(parsed, dict):
                return cleaned, parsed
            last_error = f"Parsed YAML is {type(parsed).__name__}, not dict"
        except yaml.YAMLError as e:
            last_error = str(e)
            print(f"  ⚠ Attempt {i+1}/{MAX_RETRIES} produced invalid YAML: {last_error[:120]}")
    raise ValueError(f"LLM failed to produce valid YAML after {MAX_RETRIES} attempts. Last error: {last_error}")

yaml_text, parsed = _generate_yaml()
print(yaml_text)
required_keys = [
    "request_id", "source_type", "report_name", "grain", "metrics",
    "dimensions", "filters", "bronze_sources", "target_table", "open_questions"
]
missing = [k for k in required_keys if k not in parsed]
assert not missing, f"Missing required keys: {missing}"
assert parsed["request_id"] == TICKET_KEY, f"request_id mismatch: {parsed['request_id']}"
assert parsed["source_type"] == "jira", f"source_type mismatch: {parsed['source_type']}"

# target_table and bronze_sources may legitimately be empty for sparse tickets
target = parsed.get("target_table")
if target:
    assert str(target).startswith(f"{CATALOG}.{GOLD_SCHEMA}."), (
        f"target_table must start with {CATALOG}.{GOLD_SCHEMA}. — got: {target}"
    )
else:
    # Ensure open_questions flags the missing target
    if not any("target" in str(q).lower() or "gold" in str(q).lower()
               for q in (parsed.get("open_questions") or [])):
        parsed.setdefault("open_questions", []).append(
            "What should the gold target table be named?"
        )

bronze_refs = parsed.get("bronze_sources") or []
invalid = [s for s in bronze_refs if s not in existing_bronze_tables]
assert not invalid, f"YAML referenced nonexistent bronze tables: {invalid}"

# Re-serialise the (possibly patched) YAML
yaml_text = yaml.dump(parsed, default_flow_style=False, sort_keys=False)
print("\n--- Final validated spec ---")
print(yaml_text)
print("✓ YAML validation passed.")

# COMMAND ----------

# DBTITLE 1,Persist Spec to agentops.control.intake_specs
# ---------------------------------------------------------------------------
# Persist the YAML spec to agentops.control.intake_specs (Delta table)
# Stage 02 reads from this table to pick up new specs.
# ---------------------------------------------------------------------------
from datetime import datetime, timezone
from pyspark.sql.types import StructType, StructField, StringType, TimestampType, ArrayType

# Ensure schema exists
spark.sql(f"CREATE SCHEMA IF NOT EXISTS {CATALOG}.control")

spec_schema = StructType([
    StructField("request_id", StringType(), False),
    StructField("source_type", StringType()),
    StructField("report_name", StringType()),
    StructField("grain", StringType()),
    StructField("target_table", StringType()),
    StructField("yaml_spec", StringType()),
    StructField("ticket_summary", StringType()),
    StructField("ticket_status", StringType()),
    StructField("child_ticket_keys", StringType()),
    StructField("bronze_sources", StringType()),
    StructField("processed_at", TimestampType()),
])

child_keys = ",".join(ct["key"] for ct in child_tickets) if child_tickets else None
bronze_csv  = ",".join(parsed.get("bronze_sources") or []) or None

result_row = [(
    TICKET_KEY,
    "jira",
    parsed.get("report_name"),
    str(parsed.get("grain")) if parsed.get("grain") else None,
    str(parsed.get("target_table")) if parsed.get("target_table") else None,
    yaml_text,
    ticket_summary,
    ticket_status,
    child_keys,
    bronze_csv,
    datetime.now(timezone.utc),
)]

result_df = spark.createDataFrame(result_row, schema=spec_schema)

# Append (or create) the Delta table
result_df.write.format("delta").mode("append").option(
    "mergeSchema", "true"
).saveAsTable(f"{CATALOG}.control.intake_specs")

print(f"\u2713 Spec persisted to {CATALOG}.control.intake_specs")
display(spark.table(f"{CATALOG}.control.intake_specs").where(f"request_id = '{TICKET_KEY}'"))